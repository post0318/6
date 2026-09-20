"""FG_Trading_H_I.xlsm 생성: openpyxl로 시트 뼈대를 만들고, COM으로 VBA(.bas) 주입 + 버튼 추가.

VBA 프로젝트 개체 모델 접근(AccessVBOM)은 빌드 동안만 임시로 켜고 finally에서 원래대로(값 삭제) 되돌린다.
"""

from __future__ import annotations

import sys
import winreg
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
XLSX = HERE / "_skeleton.xlsx"
XLSM = HERE / "FG_Trading_H_I.xlsm"
BAS = HERE / "FG_Trading.bas"

HDR = PatternFill("solid", fgColor="1F3A5F")
HFONT = Font(bold=True, color="FFFFFF")
INPUT = PatternFill("solid", fgColor="FFF4CC")

LABELS = [
    ("DEPOSIT", "입금(신규 트랜치)"), ("WITHDRAW", "출금"), ("RAMP_BUY", "초기 분할편입 매수"),
    ("WEEKLY_BUY", "정기 추가매수"), ("CRASH_BUY", "급락 매수"), ("SELL", "탐욕 매도"),
    ("WD_SELL", "출금 충당 매도"), ("FILL_ADJ", "실제 체결 보정"),
    ("RAMP", "초기 분할편입 중"), ("ACTIVE", "정기매수 진행(초기편입 완료)"), ("SOLD", "매도 후 재진입 대기"),
    ("IDLE", "매수 비활성(FG 매도기준 이상)"), ("CLOSED", "종료"),
    ("MSG_NODATE", "주문일자가 비어 있습니다. Settings 시트 L2 셀(OrderDate)에 오늘 주문일자를 입력한 뒤 다시 실행하세요."),
    ("MSG_DUP", "이미 처리한 날짜입니다. OrderDate는 마지막 처리일(LastRunDate, L3)보다 뒤여야 합니다."),
    ("MSG_NOMKT", "MARKET 시트가 비어 있습니다. 일자 / FG / ETF 종가를 먼저 입력하세요."),
    ("MSG_MKTDATE", "MARKET 마지막 행의 일자는 OrderDate보다 이전이어야 합니다. (직전 영업일 행까지만 입력)"),
    ("MSG_PX", "MARKET 마지막 행의 ETF 종가가 올바르지 않습니다."),
    ("MSG_DONE", "처리 완료 ({DATE}) - 기준 FG {FG}, 기준가 {PX}\n입금 {DEP}건 / 출금 {WD}건 처리, 오류 {ERR}건, 미래일자 대기 {FUT}건\n관리 중인 입금건(트랜치) {TR}개 / 생성된 주문 {ORD}행 (ORDERS 시트 확인)"),
    ("MSG_NOINV", "※ 관리 중인 투자자가 없습니다. INVESTOR 시트에 투자자ID / 일자 / 금액 / 전략(H 또는 I)을 붙여넣고 다시 실행하세요."),
    ("MSG_ERRROWS", "※ INVESTOR 시트 처리상태 열에 ERR 표시된 행을 확인하세요(일자·금액·전략 오류). 수정 후 처리상태를 비우면 다음 실행 때 다시 처리됩니다."),
    ("MSG_SMALL", "※ 정수 주 단위(INT)라 1주 가격에 못 미쳐 건너뛴 매수가 {SKIP}건 있습니다. 금액이 작은 투자자는 소수 단위(FRAC) 또는 ETF 단가를 확인하세요."),
    ("MSG_NOORDER", "※ 오늘은 신규 매매 대상이 없어 주문이 0행입니다. 진행 상황은 STATE 시트에서 확인하세요."),
]

GUIDE = [
    "FG 전략 H/I 일일 주문 산출 워크북 (운영용)",
    "",
    "[전략] H = 초기 40% 10영업일 분할편입, 21영업일마다 총자산 10% 추가매수, FG(전일)>=79 전량 매도, FG<31 급락 시 잔여 현금 전량 매수.",
    "       I = H와 동일하나 FG>=79 매도를 국면당 1회 50%만 수행.",
    "[신호] 직전 미국 영업일(MARKET 마지막 행)의 FG로 판단. 수량 산정 기준가는 같은 행의 ETF 종가(전일 종가 추정).",
    "",
    "[매일 절차]",
    "1) MARKET 시트 맨 아래에 직전 미국 영업일 행 추가: 날짜 / FG / ETF 종가 (FG는 최근값만 있으면 됨).",
    "2) 내부 시스템에서 내려받은 오늘 자 투자자 입출금을 INVESTOR 시트 맨 아래에 붙여넣기 (열: 투자자ID, 일자, 금액, 전략). 출금은 음수. 마지막 열(처리상태)은 비워 둠.",
    "3) Settings 시트의 OrderDate(주문일자)를 오늘 날짜로 입력 후 [주문 산출] 버튼(RunDay) 실행.",
    "4) ORDERS 시트에서 투자자별 매수수량/매도수량 확인. 한 행에서 둘 중 하나만 0이 아님(동시 발생 불가, 순매매 기준).",
    "5) 체결 후 실제 수량(H열)과 체결가(I열)를 입력하고 [체결 반영] 버튼(ApplyFills) 실행 - 추정 대비 차이가 STATE에 보정됨.",
    "",
    "[핵심 확인] STATE 시트(투자자·입금건별 1행): 현재 단계(초기 분할편입 n/10 또는 정기매수 진행), 초기편입 누적목표비중, 다음 정기매수까지 남은 영업일과 예정금액·비중증가,",
    "       매도신호가 뜨면 목표비중(H=0%, I=현재 보유의 50%)과 매도수량. ORDERS의 '주문 사유'에는 오늘 주문이 몇 회차 초기편입/정기매수/급락매수/매도 때문인지 표시됨.",
    "[시트] INVESTOR=입출금 DB(누적, 삭제 금지) / STATE=트랜치별 현재 상태(매 실행 덮어씀) / LOG=매매·입출금·보정 이벤트만 누적 / ORDERS=일별 주문서.",
    "[규칙] 입금 1건 = 트랜치 1개(각자 입금일부터 분할편입). 출금은 현금 우선, 부족분은 보유수량 비례 매도. RunDay는 영업일당 정확히 1회.",
    "[개인정보] 이 파일은 내부망 전용. 이름 등 식별정보 대신 투자자 ID만 사용할 것.",
]


def build_skeleton() -> None:
    wb = Workbook()
    g = wb.active
    g.title = "Guide"
    for i, line in enumerate(GUIDE, 1):
        g.cell(i, 1, line)
    g["A1"].font = Font(bold=True, size=14)
    g.column_dimensions["A"].width = 150

    s = wb.create_sheet("Settings")
    heads = ["전략", "초기비중", "분할일수", "추가매수주기(일)", "추가매수비중", "매도기준FG", "급락기준FG", "급락매수비율", "매도비율"]
    for c, h in enumerate(heads, 1):
        s.cell(1, c, h)
    s.append(["H", 0.40, 10, 21, 0.10, 79, 31, 1.0, 1.0])
    s.append(["I", 0.40, 10, 21, 0.10, 79, 31, 1.0, 0.5])
    for c, h in enumerate(["key", "설명", "값"], 10):
        s.cell(1, c, h)
    opts = [
        ("OrderDate", "주문일자(오늘)", None),
        ("LastRunDate", "마지막 처리일(자동)", None),
        ("QtyMode", "수량 단위 INT=정수주/FRAC=소수", "INT"),
        ("Ticker", "매매 ETF(참고)", ""),
        ("Silent", "메시지창 끄기 Y/N (자동화용)", "N"),
        ("LastMessage", "마지막 처리 메시지(자동)", ""),
    ]
    for i, (k, d, v) in enumerate(opts, 2):
        s.cell(i, 10, k)
        s.cell(i, 11, d)
        s.cell(i, 12, v)
        s.cell(i, 12).fill = INPUT
    s["L2"].number_format = "yyyy-mm-dd"
    s["L3"].number_format = "yyyy-mm-dd"
    dv = DataValidation(type="list", formula1='"INT,FRAC"', allow_blank=False)
    s.add_data_validation(dv)
    dv.add("L4")
    s.cell(1, 14, "이벤트")
    s.cell(1, 15, "표시명")
    for i, (k, v) in enumerate(LABELS, 2):
        s.cell(i, 14, k)
        s.cell(i, 15, v)
    s["A5"] = "▶ 매일 [주문 산출] 실행 전: L2 셀(OrderDate)에 오늘 주문일자를 입력하세요. (노란색 칸)"
    s["A5"].font = Font(bold=True, color="C00000")
    for c in list(range(1, 10)) + [10, 11, 12, 14, 15]:
        s.cell(1, c).fill, s.cell(1, c).font = HDR, HFONT
    for col, w in zip("ABCDEFGHIJKLMNO", [8, 10, 10, 16, 12, 12, 12, 12, 10, 14, 34, 14, 3, 14, 22]):
        s.column_dimensions[col].width = w

    def sheet(name: str, heads: list[str], widths: list[int]):
        ws = wb.create_sheet(name)
        for c, h in enumerate(heads, 1):
            cell = ws.cell(1, c, h)
            cell.fill, cell.font = HDR, HFONT
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            ws.column_dimensions[cell.column_letter].width = widths[c - 1]
        ws.freeze_panes = "A2"
        return ws

    m = sheet("MARKET", ["일자", "FG(공포탐욕)", "ETF 종가"], [14, 14, 14])
    m.column_dimensions["A"].number_format = "yyyy-mm-dd"
    i = sheet("INVESTOR", ["투자자ID", "일자", "금액(입금+/출금-)", "전략(H/I)", "처리상태"], [14, 14, 20, 12, 14])
    dv2 = DataValidation(type="list", formula1='"H,I"', allow_blank=True)
    i.add_data_validation(dv2)
    dv2.add("D2:D50000")
    stt = sheet("STATE", ["투자자ID", "트랜치", "전략", "시작일", "현금", "보유수량", "매수활성", "매도완료", "정기매수경과일", "경과영업일", "종료", "평가액", "주식비중",
                          "단계코드", "현재 단계", "초기편입 완료회차", "초기편입 누적목표비중", "다음 정기매수까지 영업일", "다음 정기매수 예정금액", "예정 비중증가(%p)",
                          "매도신호 시 목표비중", "매도신호 시 매도수량"],
                [12, 8, 8, 12, 16, 12, 10, 10, 14, 10, 8, 16, 10, 10, 30, 12, 14, 14, 16, 12, 14, 14])
    for col in "MQTU":
        stt.column_dimensions[col].number_format = "0.0%"
    lg = sheet("LOG", ["일자", "투자자ID", "트랜치", "이벤트", "구분", "수량(+매수/-매도)", "단가", "금액", "주식비중", "상태", "현금증감/입출금", "초기편입 회차", "비중변화(%p)"],
               [12, 12, 8, 12, 20, 16, 12, 16, 10, 12, 16, 12, 12])
    lg.column_dimensions["I"].number_format = "0.0%"
    lg.column_dimensions["M"].number_format = "0.0%"
    sheet("ORDERS", ["주문일자", "투자자ID", "매수수량", "매도수량", "추정기준가(전일종가)", "추정금액(+매수/-매도)", "점검", "실제체결수량", "실제체결가", "반영상태", "주문 사유", "주문후 투자비중"],
          [12, 12, 12, 12, 18, 20, 8, 14, 12, 14, 44, 14])
    wb.save(XLSX)


def set_vbom(on: bool, prev) -> None:
    key = winreg.CreateKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Office\16.0\Excel\Security")
    with key:
        if on:
            winreg.SetValueEx(key, "AccessVBOM", 0, winreg.REG_DWORD, 1)
        elif prev is None:
            try:
                winreg.DeleteValue(key, "AccessVBOM")
            except FileNotFoundError:
                pass
        else:
            winreg.SetValueEx(key, "AccessVBOM", 0, winreg.REG_DWORD, prev)


def read_vbom():
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Office\16.0\Excel\Security") as k:
            return winreg.QueryValueEx(k, "AccessVBOM")[0]
    except FileNotFoundError:
        return None


def inject() -> None:
    import win32com.client as wc

    build_skeleton()
    xl = wc.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(str(XLSX))
        tmp = HERE / "_FG_Trading_crlf.bas"
        tmp.write_bytes(BAS.read_bytes().replace(b"\r\n", b"\n").replace(b"\n", b"\r\n"))
        wb.VBProject.VBComponents.Import(str(tmp))
        tmp.unlink()
        ws = wb.Worksheets("Settings")
        for i, (cap, macro) in enumerate([("주문 산출 (RunDay)", "RunDay"), ("체결 반영 (ApplyFills)", "ApplyFills")]):
            left, top = 20 + i * 190, 130
            b = ws.Buttons().Add(left, top, 180, 36)
            b.OnAction = macro
            b.Caption = cap
        wb.SaveAs(str(XLSM), 52)
        wb.Close(False)
    finally:
        xl.Quit()


if __name__ == "__main__":
    prev = read_vbom()
    print("AccessVBOM before:", prev)
    set_vbom(True, prev)
    try:
        inject()
        print("built", XLSM)
    finally:
        set_vbom(False, prev)
        print("AccessVBOM restored to:", read_vbom())
    XLSX.unlink(missing_ok=True)
    sys.exit(0)
